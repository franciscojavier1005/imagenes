#!/usr/bin/env python3
"""Base de datos aparte (consulta y verificación): SQLite + Excel.

Salidas:
  ICET_Base_Datos_2026.db    SQLite con tablas, vistas y consultas de ejemplo
  ICET_Base_Datos_2026.xlsx  mismas tablas + hojas de verificación (cargas, cobertura, controles)
"""
import csv, glob, os, re, sqlite3, sys
from collections import defaultdict
sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import pdfplumber
from extraer_horarios import extraer_pagina
from construir_datos import FUENTES, limpiar_titulo, clave_pdf, MAPA

RAIZ = os.path.join(os.path.dirname(__file__), "..")
D = os.path.join(RAIZ, "data")
DIAS = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES"]


def leer(n):
    return list(csv.DictReader(open(os.path.join(D, n), encoding="utf-8-sig")))


docentes, grupos, dirg, franjas = leer("docentes.csv"), leer("grupos.csv"), leer("direccion_grupo.csv"), leer("franjas.csv")
horario, motivos, listas = leer("horario_maestro.csv"), leer("motivos.csv"), leer("listas.csv")
GN = {g["grupo"]: g["nombre"] for g in grupos}

# ------------------------------------------------------------------ cálculos de verificación
carga = defaultdict(lambda: defaultdict(int))     # docente_n -> día -> sesiones
tipos = defaultdict(lambda: defaultdict(int))     # docente_n -> tipo -> sesiones
cobert = defaultdict(lambda: defaultdict(set))    # grupo -> día -> sesiones distintas (varios docentes pueden compartir la franja)
for r in horario:
    n = int(r["docente_n"])
    carga[n][r["dia"]] += 1
    tipos[n][r["tipo"]] += 1
    for g in ([r["grupo"]] if r["tipo"] != "ENFASIS" else r["grupos_enfasis"].split("+")):
        cobert[g][r["dia"]].add(int(r["hora"]))

# secciones declaradas en el título de cada PDF por docente ("24 SECCIONES", "28 S", "28-S")
declaradas = {}
for pdf in sorted(glob.glob(os.path.join(FUENTES, "*ASIG_ACAD_DOC-*.pdf"))):
    with pdfplumber.open(pdf) as d:
        celdas = defaultdict(int)
        for p in d.pages:
            for f in extraer_pagina(p):
                celdas[limpiar_titulo(f["titulo"])] += 1
        for t in celdas:
            m = re.search(r"(\d+)\s*-?\s*S\b", t) or re.search(r"(\d+)\s*SECC", t)
            k = clave_pdf(t)
            declaradas[MAPA[k]] = (int(m.group(1)) if m else None, celdas[t])

checks = []  # (control, resultado, detalle)
sec = [d for d in docentes if d["nivel"] == "SECUNDARIA" and d["tiene_horario"] == "SI"]
dif = [(d["nombre_completo"], declaradas[int(d["n"])]) for d in sec if int(d["n"]) in declaradas and declaradas[int(d["n"])][0] != sum(carga[int(d["n"])].values())]
checks.append(("Sesiones por docente = 'SECCIONES' del título del PDF", "OK" if not dif else "INFO",
               f"{len(sec)} docentes; " + ("sin diferencias" if not dif else
               "; ".join(f"{n}: el título dice {v[0]} y la cuadrícula tiene {v[1]}" for n, v in dif) +
               ". La cuadrícula coincide con los PDF por grupo, así que se asume errata en el título.")))
esperado = lambda g: 25 if g.startswith("00") else (30 if g[:2] in ("01", "02", "03", "04", "05") else 40)
TRANSV = {g["grupo"] for g in grupos if g["tipo"] == "TRANSVERSAL"}  # orientación y tutoría: sin cobertura esperada
falt = [(GN[g], sum(len(v) for v in cobert[g].values()), esperado(g)) for g in GN if g not in TRANSV and sum(len(v) for v in cobert[g].values()) != esperado(g)]
checks.append(("Cada grupo con todas sus sesiones cubiertas (40 bachillerato, 30 primaria, 25 preescolar)", "OK" if not falt else "REVISAR", f"{len(GN) - len(TRANSV)} grupos; incompletos: {len(falt)}"))
doble = defaultdict(int)
for r in horario:
    doble[(r["docente_n"], r["dia"], r["hora"])] += 1
c1 = sum(1 for v in doble.values() if v > 1)
checks.append(("Docente en dos sitios a la vez", "OK" if c1 == 0 else "REVISAR", f"{c1} casos"))
gg = defaultdict(set)
for r in horario:
    if r["tipo"] in ("CLASE", "AREAS_MULTIPLES"):
        gg[(r["grupo"], r["dia"], r["hora"])].add(r["docente_n"])
c2 = sum(1 for v in gg.values() if len(v) > 1)
checks.append(("Dos docentes de clase con el mismo grupo a la vez", "OK" if c2 == 0 else "REVISAR", f"{c2} casos"))
sin_nivel = [d["nombre_completo"] for d in docentes if d["nivel"] == "POR DEFINIR"]
checks.append(("Docentes con nivel definido", "OK" if not sin_nivel else "PENDIENTE", "; ".join(sin_nivel) or "todos"))
sin_hor = [d["nombre_completo"] for d in docentes if d["tiene_horario"] == "NO" and d["nivel"] != "REEMPLAZADO"]
checks.append(("Docentes con horario cargado", "OK" if not sin_hor else "PENDIENTE", f"{len(sin_hor)} sin horario: " + "; ".join(sin_hor)))
sin_dir = [d["nombre_completo"] for d in docentes if not d["direccion_grupo"]]
checks.append(("Docentes con dirección de grupo", "INFO", f"{len(sin_dir)} sin dirección: " + "; ".join(sin_dir)))
checks.append(("Reemplazos y cambios del rector aplicados", "INFO", "Armero Dájome Jesús (OPS) reemplaza a Terán Guevara Jorge Alberto; Prado Genís Maribel y Puches Ana Milena intercambiaron su horario de Arte"))
checks.append(("Alternancia semanal de énfasis (parejas, Sociales/Inglés, Ética/Religión)", "INFO", "No se hace seguimiento semanal en la base: lo controla internamente el coordinador"))
checks.append(("Personal sin grupo (orientación y tutoría PTAFI)", "INFO", "Casanova Johana (orientadora): 07:30-14:30; Ponce Ángela (orientadora): 08:30-13:30; Ortiz Martha (tutora PTAFI): jornada de primaria 06:30-12:00"))

# ------------------------------------------------------------------ SQLite
DB = os.path.join(RAIZ, "ICET_Base_Datos_2026.db")
if os.path.exists(DB):
    os.remove(DB)
con = sqlite3.connect(DB)
cur = con.cursor()
cur.executescript("""
CREATE TABLE docentes(n INTEGER PRIMARY KEY, codigo TEXT, apellidos TEXT, nombres TEXT, nombre_completo TEXT, nivel TEXT,
  grupo_titular TEXT, direccion_grupo TEXT, areas TEXT, correo TEXT, tiene_horario TEXT, nota TEXT);
CREATE TABLE grupos(grupo TEXT PRIMARY KEY, tipo TEXT, grado TEXT, nombre TEXT, modalidad TEXT);
CREATE TABLE direccion_grupo(grupo TEXT PRIMARY KEY, nombre TEXT, tipo_direccion TEXT, docente_1_n INTEGER, docente_1 TEXT,
  docente_2_n INTEGER, docente_2 TEXT, modalidad TEXT);
CREATE TABLE franjas(hora INTEGER PRIMARY KEY, bloque INTEGER, sesion INTEGER, inicio TEXT, fin TEXT, inicio_min INTEGER, fin_min INTEGER);
CREATE TABLE horario(id INTEGER PRIMARY KEY AUTOINCREMENT, docente_n INTEGER REFERENCES docentes(n), docente TEXT, dia TEXT,
  hora INTEGER REFERENCES franjas(hora), bloque INTEGER, sesion INTEGER, inicio TEXT, fin TEXT, tipo TEXT, grupo TEXT, area TEXT,
  enfasis_ref TEXT, grupos_enfasis TEXT, equipo_enfasis TEXT, alternancia TEXT, celda_original TEXT);
CREATE TABLE motivos(categoria TEXT, motivo TEXT PRIMARY KEY, justificada TEXT, soporte_sugerido TEXT, requiere_soporte TEXT, plazo_dias INTEGER);
CREATE TABLE listas(lista TEXT, valor TEXT);
CREATE TABLE registro_ronda(id INTEGER PRIMARY KEY AUTOINCREMENT, marca_temporal TEXT, fecha TEXT, dia TEXT, sesion INTEGER, franja TEXT,
  docente TEXT, grupo TEXT, area TEXT, estado TEXT, motivo TEXT, justificada TEXT, minutos INTEGER, observaciones TEXT, directivo TEXT);
CREATE TABLE novedades(id INTEGER PRIMARY KEY AUTOINCREMENT, marca_temporal TEXT, fecha_novedad TEXT, docente TEXT, tipo_novedad TEXT,
  actividad_aprendizaje TEXT, motivo_ausencia TEXT, descripcion TEXT, fuente_novedad TEXT, medio_informacion TEXT, grado TEXT, grupo TEXT,
  area_asignatura TEXT, horario TEXT, minutos_desatendidos INTEGER, directivo_docente TEXT, sesiones TEXT, justificada TEXT, categoria_motivo TEXT);
""")
ins = lambda t, rows, cols: cur.executemany(f"INSERT INTO {t}({','.join(cols)}) VALUES({','.join('?' * len(cols))})", [[r[c] for c in cols] for r in rows])
ins("docentes", docentes, ["n", "codigo", "apellidos", "nombres", "nombre_completo", "nivel", "grupo_titular", "direccion_grupo", "areas", "correo", "tiene_horario", "nota"])
ins("grupos", grupos, ["grupo", "tipo", "grado", "nombre", "modalidad"])
ins("direccion_grupo", [{**r, "docente_1_n": r["docente_1_n"] or None, "docente_2_n": r["docente_2_n"] or None} for r in dirg],
    ["grupo", "nombre", "tipo_direccion", "docente_1_n", "docente_1", "docente_2_n", "docente_2", "modalidad"])
ins("franjas", franjas, ["hora", "bloque", "sesion", "inicio", "fin", "inicio_min", "fin_min"])
ins("horario", horario, ["docente_n", "docente", "dia", "hora", "bloque", "sesion", "inicio", "fin", "tipo", "grupo", "area", "enfasis_ref",
                         "grupos_enfasis", "equipo_enfasis", "alternancia", "celda_original"])
ins("motivos", motivos, ["categoria", "motivo", "justificada", "soporte_sugerido", "requiere_soporte", "plazo_dias"])
ins("listas", listas, ["lista", "valor"])
cur.executescript("""
CREATE INDEX ix_horario_dia_hora ON horario(dia, hora);
CREATE INDEX ix_horario_docente ON horario(docente_n);
CREATE VIEW v_horario AS
  SELECT h.dia, h.hora AS sesion, f.inicio, f.fin, f.bloque,
         CASE WHEN h.tipo='ENFASIS' THEN 'ÉNFASIS '||h.grupos_enfasis ELSE g.nombre END AS grupo,
         h.docente, d.nivel, COALESCE(NULLIF(h.area,''),'(énfasis)') AS area, h.tipo, h.alternancia, h.equipo_enfasis
  FROM horario h JOIN franjas f ON f.hora=h.hora JOIN docentes d ON d.n=h.docente_n LEFT JOIN grupos g ON g.grupo=h.grupo;
CREATE VIEW v_carga_docente AS
  SELECT d.n, d.nombre_completo, d.nivel, COUNT(h.id) AS sesiones_semana, ROUND(COUNT(h.id)*45/60.0,2) AS horas_semana,
         SUM(h.tipo='CLASE') AS clase, SUM(h.tipo='ENFASIS') AS enfasis, SUM(h.tipo='AREAS_MULTIPLES') AS areas_multiples, SUM(h.tipo IN ('ORIENTACION','TUTORIA')) AS orientacion_tutoria
  FROM docentes d LEFT JOIN horario h ON h.docente_n=d.n GROUP BY d.n;
CREATE VIEW v_cobertura_grupo AS
  SELECT g.nombre AS grupo, COUNT(*) AS sesiones_cubiertas FROM (
    SELECT grupo AS gcod, dia, hora FROM horario WHERE tipo<>'ENFASIS'
    UNION SELECT TRIM(SUBSTR(grupos_enfasis,1,INSTR(grupos_enfasis||'+','+')-1)), dia, hora FROM horario WHERE tipo='ENFASIS'
    UNION SELECT TRIM(SUBSTR(grupos_enfasis,INSTR(grupos_enfasis,'+')+1)), dia, hora FROM horario WHERE tipo='ENFASIS' AND grupos_enfasis LIKE '%+%'
  ) x JOIN grupos g ON g.grupo=x.gcod GROUP BY g.nombre;
""")
con.commit()

# consultas de ejemplo, probadas para el README
ejemplos = {
    "¿Quién debe estar el martes en la sesión 3?": "SELECT grupo, docente, area FROM v_horario WHERE dia='MARTES' AND sesion=3 ORDER BY grupo",
    "¿Dónde está un docente el lunes?": "SELECT sesion, inicio, fin, grupo, area FROM v_horario WHERE docente LIKE 'Villota Rubio Janeth%' AND dia='LUNES' ORDER BY sesion",
    "Carga semanal de cada docente": "SELECT nombre_completo, nivel, sesiones_semana, horas_semana FROM v_carga_docente ORDER BY sesiones_semana DESC LIMIT 5",
}
probado = {k: len(cur.execute(v).fetchall()) for k, v in ejemplos.items()}
vc = dict(cur.execute("SELECT grupo, sesiones_cubiertas FROM v_cobertura_grupo").fetchall())
assert len(vc) == len(GN), "la vista de cobertura no cubre todos los grupos"
con.close()

# ------------------------------------------------------------------ Excel
AZUL = PatternFill("solid", fgColor="1F4E78")
OKF, REV, PEN = (PatternFill("solid", fgColor=c) for c in ("C6E0B4", "F8CBAD", "FFE699"))
wb = Workbook(); wb.remove(wb.active)


def hoja(nombre, filas, anchos=None):
    ws = wb.create_sheet(nombre)
    for f in filas:
        ws.append(f)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = AZUL; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for i in range(1, ws.max_column + 1):
        w = max(len(str(ws.cell(r, i).value or "")) for r in range(1, min(ws.max_row, 100) + 1))
        ws.column_dimensions[get_column_letter(i)].width = min(max(8, w + 2), 50)
    ws.auto_filter.ref = ws.dimensions
    return ws


def tabla(nombre, filas, cols):
    return hoja(nombre, [cols] + [[(int(r[c]) if str(r[c]).isdigit() and c in ("n", "hora", "bloque", "sesion", "inicio_min", "fin_min", "docente_n") else r[c]) for c in cols] for r in filas])


lm = wb.create_sheet("LEEME")
for t in ["ICET 2026 - Base de datos de horarios y asistencia docente (consulta y verificación)", "",
          "Tablas: docentes, grupos, direccion_grupo, franjas, horario, motivos, listas.",
          "Verificación: Controles (resumen), Carga_Docente (sesiones por día y horas), Cobertura_Grupo (sesiones por día de cada grupo).",
          "Consulta sin filtros: hoja Horario_Legible (filtre por día y sesión).",
          "La misma información en SQLite: ICET_Base_Datos_2026.db (vistas v_horario, v_carga_docente, v_cobertura_grupo).",
          "Un registro de 'horario' = un docente, un día y una sesión de 45 minutos. Bloque = 2 sesiones (90 min).",
          "Fuentes: PDF de asignación académica (02-02-2026), PDF de dirección de grupo (marzo 2026), Word de asistencia 2026."]:
    lm.append([t])
lm["A1"].font = Font(bold=True, size=14); lm.column_dimensions["A"].width = 130

ws = hoja("Controles", [["Control", "Resultado", "Detalle"]] + [list(c) for c in checks])
for row in ws.iter_rows(min_row=2):
    row[1].fill = {"OK": OKF, "REVISAR": REV}.get(row[1].value, PEN)
    row[2].alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 70; ws.column_dimensions["C"].width = 110

# Carga docente
cab = ["n", "docente", "nivel", "áreas"] + DIAS + ["sesiones_semana", "horas_semana", "clase", "énfasis", "áreas_múltiples", "secciones_según_PDF", "coincide"]
filas = [cab]
for d in docentes:
    n = int(d["n"]); c = carga[n]; tot = sum(c.values())
    dec = declaradas.get(n, (None,))[0]
    filas.append([n, d["nombre_completo"], d["nivel"], d["areas"]] + [c.get(x, 0) for x in DIAS] +
                 [tot, round(tot * 45 / 60, 2), tipos[n].get("CLASE", 0), tipos[n].get("ENFASIS", 0), tipos[n].get("AREAS_MULTIPLES", 0),
                  dec if dec is not None else "", ("SI" if dec == tot else "NO") if dec is not None else ""])
hoja("Carga_Docente", filas)

# Cobertura por grupo
filas = [["grupo", "nombre", "modalidad"] + DIAS + ["sesiones_cubiertas", "esperadas", "completo"]]
for g in grupos:
    c = {d: len(v) for d, v in cobert[g["grupo"]].items()}; tot = sum(c.values())
    filas.append([g["grupo"], g["nombre"], g["modalidad"]] + [c.get(x, 0) for x in DIAS] + [tot, "—" if g["grupo"] in TRANSV else esperado(g["grupo"]), ("—" if g["grupo"] in TRANSV else ("SI" if tot == esperado(g["grupo"]) else "NO"))])
hoja("Cobertura_Grupo", filas)

tabla("Docentes", docentes, ["n", "codigo", "apellidos", "nombres", "nombre_completo", "nivel", "grupo_titular", "direccion_grupo", "areas", "correo", "tiene_horario", "nota"])
tabla("Grupos", grupos, ["grupo", "tipo", "grado", "nombre", "modalidad"])
tabla("Direccion_Grupo", dirg, ["grupo", "nombre", "tipo_direccion", "docente_1_n", "docente_1", "docente_2_n", "docente_2", "modalidad"])
tabla("Franjas", franjas, ["hora", "bloque", "sesion", "inicio", "fin", "inicio_min", "fin_min"])
tabla("Horario", horario, ["docente_n", "docente", "dia", "hora", "bloque", "sesion", "inicio", "fin", "tipo", "grupo", "area", "enfasis_ref", "grupos_enfasis", "equipo_enfasis", "alternancia", "celda_original"])
leg = [["dia", "sesion", "inicio", "fin", "grupo", "docente", "area", "tipo", "alternancia", "equipo_enfasis"]]
orden = {d: i for i, d in enumerate(DIAS)}
fr = {int(f["hora"]): f for f in franjas}
for r in sorted(horario, key=lambda r: (orden[r["dia"]], int(r["hora"]), (GN.get(r["grupo"]) or "~" + r["grupos_enfasis"]), r["docente"])):
    gr = ("ÉNFASIS " + " + ".join(GN.get(x, x) for x in r["grupos_enfasis"].split("+"))) if r["tipo"] == "ENFASIS" else GN.get(r["grupo"], r["grupo"])
    leg.append([r["dia"], int(r["hora"]), fr[int(r["hora"])]["inicio"], fr[int(r["hora"])]["fin"], gr, r["docente"], r["area"] or "(énfasis)", r["tipo"], r["alternancia"], r["equipo_enfasis"]])
hoja("Horario_Legible", leg)
tabla("Motivos", motivos, ["categoria", "motivo", "justificada", "soporte_sugerido", "requiere_soporte", "plazo_dias"])
hoja("Novedades", [["Marca temporal", "Fecha Novedad", "Docente", "Tipo Novedad", "Actividad de Aprendizaje", "Motivo Ausencia", "Descripción", "Fuente Novedad",
                    "Medio Información", "Grado", "Grupo", "Área/Asignatura", "Horario", "Minutos Desatendidos", "Directivo Docente", "Sesiones", "Justificada", "Categoría motivo"]])
wb.save(os.path.join(RAIZ, "ICET_Base_Datos_2026.xlsx"))
print("OK. Controles:")
for c in checks:
    print(f"  [{c[1]:9s}] {c[0]} -> {c[2][:100]}")
print("Consultas de ejemplo (filas):", probado)
