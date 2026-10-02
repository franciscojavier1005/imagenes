#!/usr/bin/env python3
"""ICET_Ubicacion_Docentes_Bachillerato.xlsx: ¿dónde está cada docente de bachillerato?

Hojas: Ubicacion (por grupo), Por_Docente (A-Z), Horario_Bach (datos), Franjas, Grupos, LEEME.
Al abrir, la hoja Ubicacion aplica sola el filtro con la fecha y la hora del dispositivo. Los desplegables de fecha, bloque y sesión
permiten consultar otro momento. Solo bachillerato: clases y énfasis (sin preescolar ni primaria).
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

RAIZ = os.path.join(os.path.dirname(__file__), "..")
D = os.path.join(RAIZ, "data")
OUT = os.path.join(RAIZ, "ICET_Ubicacion_Docentes_Bachillerato.xlsx")
rd = lambda n: list(csv.DictReader(open(os.path.join(D, n), encoding="utf-8-sig")))
AZUL = PatternFill("solid", fgColor="1F4E78")
AMARILLO = PatternFill("solid", fgColor="FFF2CC")
MAXF, NF = 1200, 60

grupos = {g["grupo"]: g for g in rd("grupos.csv")}
H = [r for r in rd("horario_maestro.csv") if r["tipo"] in ("CLASE", "ENFASIS")]
docentes = {d["n"]: d for d in rd("docentes.csv")}


def nombre_grupo(r):
    if r["tipo"] == "ENFASIS":
        return "ÉNFASIS " + " + ".join(grupos[g]["nombre"] for g in r["grupos_enfasis"].split("+"))
    return grupos[r["grupo"]]["nombre"]


def cab(ws, fila=1):
    for c in ws[fila]:
        if c.value is not None:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = AZUL; c.alignment = Alignment(wrap_text=True, vertical="center")


def anchos(ws, tope=48):
    for i in range(1, ws.max_column + 1):
        w = max(len(str(ws.cell(r, i).value or "")) for r in range(1, min(ws.max_row, 80) + 1))
        ws.column_dimensions[get_column_letter(i)].width = min(max(8, w + 2), tope)


wb = Workbook(); wb.remove(wb.active)

# ---------------------------------------------------------------- Franjas y Grupos
fr = wb.create_sheet("Franjas")
fr.append(["sesion", "bloque", "inicio", "fin", "inicio_min", "fin_min"])
for f in rd("franjas.csv"):
    fr.append([int(f["hora"]), int(f["bloque"]), f["inicio"], f["fin"], int(f["inicio_min"]), int(f["fin_min"])])
cab(fr); anchos(fr)
gr = wb.create_sheet("Grupos")
gr.append(["grupo", "nombre", "modalidad"])
for g in grupos.values():
    if g["tipo"] in ("REGULAR", "CAMINAR EN SECUNDARIA 1", "CAMINAR EN SECUNDARIA 2") and int(g["grupo"][:2] if g["grupo"][0] != "C" else 6) >= 6:
        gr.append([g["grupo"], g["nombre"], g["modalidad"]])
cab(gr); anchos(gr)

# ---------------------------------------------------------------- Horario_Bach
hz = wb.create_sheet("Horario_Bach")
hz.append(["docente_n", "docente", "dia", "sesion", "bloque", "tipo", "grupo", "area", "grupos_enfasis", "alternancia", "equipo_enfasis",
           "grupo_mostrar", "nota", "clave", "orden", "indice", "clave_docente"])
for i, r in enumerate(H, start=2):
    nota = " · ".join(x for x in (r["alternancia"], ("Con: " + r["equipo_enfasis"]) if r["equipo_enfasis"] else "") if x)
    hz.append([int(r["docente_n"]), r["docente"], r["dia"], int(r["hora"]), int(r["bloque"]), r["tipo"], r["grupo"], r["area"], r["grupos_enfasis"],
               r["alternancia"], r["equipo_enfasis"], nombre_grupo(r), nota,
               f'=C{i}&"|"&D{i}', f'=COUNTIF($N$2:N{i},N{i})', f'=N{i}&"|"&O{i}', f'=C{i}&"|"&D{i}&"|"&A{i}'])
cab(hz); hz.freeze_panes = "C2"; anchos(hz, 36); hz.auto_filter.ref = hz.dimensions

# ---------------------------------------------------------------- Ubicacion
ub = wb.create_sheet("Ubicacion")
ub["A1"] = "UBICACIÓN DE DOCENTES · BACHILLERATO · ICET 2026"
ub["A1"].font = Font(bold=True, size=15, color="1F4E78")
ub["A2"] = "Al abrir muestra sola quién está en cada grupo AHORA. Para consultar otro momento use las celdas amarillas (vacías = ahora)."
ub["A2"].font = Font(italic=True, color="52514E")
for r, t in ((3, "Fecha"), (4, "Bloque (90 min)"), (5, "Sesión (45 min)")):
    ub.cell(r, 1, t).font = Font(bold=True)
    ub.cell(r, 2).fill = AMARILLO
ub["C3"], ub["C4"], ub["C5"] = "← opcional: escriba otra fecha (dd/mm/aaaa)", "← opcional: bloque 1 a 4", "← opcional: sesión 1 a 8 (manda sobre el bloque)"
for c in ("C3", "C4", "C5"):
    ub[c].font = Font(color="74736D", italic=True)
ub["A7"], ub["A8"], ub["A9"], ub["A10"], ub["A11"] = "Consulta", "Franja", "Estado", "Docentes con clase", "Hora del dispositivo"
for r in range(7, 12):
    ub.cell(r, 1).font = Font(bold=True)
# auxiliares (columnas ocultas K:L)
ub["K1"], ub["L1"] = "fecha efectiva", '=IF(B3<>"",B3,TODAY())'
ub["K2"], ub["L2"] = "día", '=IFERROR(CHOOSE(WEEKDAY(L1,2),"LUNES","MARTES","MIERCOLES","JUEVES","VIERNES"),"FIN DE SEMANA")'
ub["K3"], ub["L3"] = "minuto", '=HOUR(NOW())*60+MINUTE(NOW())'
ub["K4"], ub["L4"] = "sesión auto", ('=IF(L3>=MAX(Franjas!$F$2:$F$9),"FUERA DE JORNADA",IFERROR(IF(L3<INDEX(Franjas!$F$2:$F$9,MATCH(L3,Franjas!$E$2:$E$9,1)),'
                                     'MATCH(L3,Franjas!$E$2:$E$9,1),"DESCANSO"),"FUERA DE JORNADA"))')
ub["K5"], ub["L5"] = "sesión manual", '=IF(B5<>"",B5,IF(B4<>"",(B4-1)*2+1,""))'
ub["K6"], ub["L6"] = "sesión efectiva", '=IF(L5<>"",L5,IF(L1=TODAY(),L4,"ELEGIR"))'
ub["K7"], ub["L7"] = "clave", '=IF(ISNUMBER(L6),L2&"|"&L6,"")'
ub["B7"] = '=TEXT(L1,"dd/mm/yyyy")&"  ·  "&L2'
ub["B8"] = '=IF(ISNUMBER(L6),INDEX(Franjas!$C$2:$C$9,L6)&" - "&INDEX(Franjas!$D$2:$D$9,L6)&"   (bloque "&INDEX(Franjas!$B$2:$B$9,L6)&", sesión "&L6&")","—")'
ub["B9"] = ('=IF(L2="FIN DE SEMANA","Fin de semana: no hay clases",IF(L6="ELEGIR","Elija un bloque o una sesión para esa fecha",'
            'IF(L6="DESCANSO","Descanso: no hay clases en este momento",IF(L6="FUERA DE JORNADA","Fuera de la jornada escolar",'
            'IF(L5<>"","Consulta manual","AHORA (automático)")))))')
ub["B10"] = '=IF(L7="",0,COUNTIF(Horario_Bach!$N$2:$N$%d,L7))' % MAXF
ub["B11"] = '=TEXT(NOW(),"HH:mm")'
for c in ("B7", "B8", "B9", "B10", "B11"):
    ub[c].font = Font(bold=True, size=12)
dvb = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=True)
dvs = DataValidation(type="list", formula1='"1,2,3,4,5,6,7,8"', allow_blank=True)
dvf = DataValidation(type="date", operator="between", formula1="DATE(2026,1,1)", formula2="DATE(2026,12,31)", allow_blank=True,
                     error="Escriba una fecha de 2026 (dd/mm/aaaa)")
for v in (dvb, dvs, dvf):
    ub.add_data_validation(v)
dvf.add("B3"); dvb.add("B4"); dvs.add("B5")
ub["B3"].number_format = "dd/mm/yyyy"
for j, t in enumerate(["#", "Grupo", "Docente", "Área", "Tipo", "Nota (alternancia / equipo del énfasis)"], start=1):
    ub.cell(13, j, t)
cab(ub, 13)
for k in range(1, NF + 1):
    f = 13 + k
    ub.cell(f, 1, k)
    ub.cell(f, 8, f'=IFERROR(MATCH($L$7&"|"&A{f},Horario_Bach!$P$2:$P${MAXF},0),"")')
    ub.cell(f, 2, f'=IF($H{f}="","",INDEX(Horario_Bach!$L$2:$L${MAXF},$H{f}))')
    ub.cell(f, 3, f'=IF($H{f}="","",INDEX(Horario_Bach!$B$2:$B${MAXF},$H{f}))')
    ub.cell(f, 4, f'=IF($H{f}="","",INDEX(Horario_Bach!$H$2:$H${MAXF},$H{f}))')
    ub.cell(f, 5, f'=IF($H{f}="","",IF(INDEX(Horario_Bach!$F$2:$F${MAXF},$H{f})="ENFASIS","Énfasis","Clase"))')
    ub.cell(f, 6, f'=IF($H{f}="","",INDEX(Horario_Bach!$M$2:$M${MAXF},$H{f}))')
ub.conditional_formatting.add(f"A14:F{13 + NF}", FormulaRule(formula=['$E14="Énfasis"'], fill=PatternFill("solid", bgColor="FFF2CC")))
for col, w in zip("ABCDEF", (18, 30, 38, 12, 10, 70)):
    ub.column_dimensions[col].width = w
ub.freeze_panes = "A14"
for c in ("H", "K", "L"):
    ub.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Por_Docente
pd_ = wb.create_sheet("Por_Docente")
pd_["A1"] = "DÓNDE ESTÁ CADA DOCENTE · BACHILLERATO"
pd_["A1"].font = Font(bold=True, size=15, color="1F4E78")
pd_["A2"] = "Usa la misma fecha, bloque y sesión de la hoja Ubicacion."
pd_["A2"].font = Font(italic=True, color="52514E")
pd_["A3"], pd_["B3"] = "Consulta", "=Ubicacion!B7"
pd_["A4"], pd_["B4"] = "Franja", "=Ubicacion!B8"
pd_["A5"], pd_["B5"] = "Estado", "=Ubicacion!B9"
for r in (3, 4, 5):
    pd_.cell(r, 1).font = Font(bold=True); pd_.cell(r, 2).font = Font(bold=True, size=12)
for j, t in enumerate(["Docente", "Dónde está (grupo)", "Área", "Tipo", "n"], start=1):
    pd_.cell(7, j, t)
cab(pd_, 7)
sec = sorted((d for d in docentes.values() if d["nivel"] == "SECUNDARIA"), key=lambda d: (d["apellidos"], d["nombres"]))
for i, d in enumerate(sec, start=8):
    pd_.cell(i, 1, d["nombre_completo"]); pd_.cell(i, 5, int(d["n"]))
    m = f'MATCH(Ubicacion!$L$2&"|"&Ubicacion!$L$6&"|"&$E{i},Horario_Bach!$Q$2:$Q${MAXF},0)'
    pd_.cell(i, 2, f'=IF(NOT(ISNUMBER(Ubicacion!$L$6)),"—",IFERROR(INDEX(Horario_Bach!$L$2:$L${MAXF},{m}),"Libre (sin clase)"))')
    pd_.cell(i, 3, f'=IF(NOT(ISNUMBER(Ubicacion!$L$6)),"",IFERROR(INDEX(Horario_Bach!$H$2:$H${MAXF},{m}),""))')
    pd_.cell(i, 4, f'=IF(NOT(ISNUMBER(Ubicacion!$L$6)),"",IFERROR(IF(INDEX(Horario_Bach!$F$2:$F${MAXF},{m})="ENFASIS","Énfasis","Clase"),""))')
ult = 7 + len(sec)
pd_.conditional_formatting.add(f"A8:D{ult}", FormulaRule(formula=['LEFT($B8,5)="Libre"'], font=Font(color="8A94A0")))
pd_.conditional_formatting.add(f"A8:D{ult}", FormulaRule(formula=['$D8="Énfasis"'], fill=PatternFill("solid", bgColor="FFF2CC")))
for col, w in zip("ABCDE", (38, 40, 14, 10, 6)):
    pd_.column_dimensions[col].width = w
pd_.column_dimensions["E"].hidden = True
pd_.freeze_panes = "A8"; pd_.auto_filter.ref = f"A7:D{ult}"

# ---------------------------------------------------------------- LEEME
lm = wb.create_sheet("LEEME")
for t in ["ICET 2026 - Ubicación de docentes de bachillerato", "",
          "1) Abra la hoja 'Ubicacion': muestra sola quién está en cada grupo en este momento (según la fecha y la hora del dispositivo).",
          "2) Para otro momento escriba la fecha y elija un bloque (1 a 4) o una sesión (1 a 8) en las celdas amarillas. Vacías = ahora.",
          "3) 'Por_Docente' lista a los docentes de A a Z y dónde está cada uno (o 'Libre') en el mismo momento.",
          "4) Se muestran solo clases y énfasis de bachillerato (6° a 11° y Caminar en Secundaria). No incluye preescolar ni primaria.",
          "5) Énfasis (fila amarilla): estudiantes del curso repartidos entre un equipo de docentes; cada uno atiende su área.",
          "",
          "Zona horaria: en Google Sheets use Archivo > Configuración > GMT-05:00 Bogotá. Excel usa la hora del dispositivo.",
          "Horario según el PDF de asignación del 02-02-2026 con los cambios del coordinador: Armero reemplaza a Terán; Prado y Puches intercambiaron Arte.",
          "Si cambia el horario, regenere con: python3 scripts/generar_ubicacion.py"]:
    lm.append([t])
lm["A1"].font = Font(bold=True, size=14); lm.column_dimensions["A"].width = 130

wb._sheets = [wb[n] for n in ("Ubicacion", "Por_Docente", "Horario_Bach", "Franjas", "Grupos", "LEEME")]
wb.active = 0
wb.save(OUT)
print("OK", OUT, os.path.getsize(OUT) // 1024, "KB;", len(H), "filas de horario;", len(sec), "docentes de bachillerato")
